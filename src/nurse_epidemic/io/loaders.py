"""Leitura de dados UPA."""

from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.worksheet.worksheet import Worksheet

from nurse_epidemic.schemas.columns import (
    COL_LINHA_EXCEL,
    COL_MES_COLETA,
    EXCEL_COL_MAP,
    EXCEL_FILL_COL_MAP,
    SHEET_COLUMN_ALIASES,
)

PROJECT_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_EXCEL_PATH = (
    PROJECT_ROOT / "data" / "raw" / "Planilha de dados- Estátistico.xlsx"
)

DROP_COLUMN_PATTERNS = ("Unnamed:",)


def _harmonize_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Unifica nomes de colunas entre abas."""
    out = df.copy()
    for alias, canonical in SHEET_COLUMN_ALIASES.items():
        if alias in out.columns and canonical not in out.columns:
            out[canonical] = out[alias]
        if alias in out.columns:
            out = out.drop(columns=[alias])
    return out


def _drop_junk_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Remove colunas vazias ou sem nome."""
    keep = [
        c
        for c in df.columns
        if str(c).strip() and not str(c).startswith(DROP_COLUMN_PATTERNS)
    ]
    return df[keep]


def _fill_colors(ws: Worksheet, n_rows: int) -> pd.DataFrame:
    """Lê a cor de preenchimento (ARGB) das colunas em EXCEL_FILL_COL_MAP.

    Args:
        ws: Aba aberta com openpyxl.
        n_rows: Número de linhas de dados lidas pelo pandas na aba.

    Returns:
        DataFrame alinhado às linhas do pandas; None quando sem cor.
    """
    headers = [cell.value for cell in ws[1]]
    colors: dict[str, list[str | None]] = {}
    for header, col in EXCEL_FILL_COL_MAP.items():
        if header not in headers:
            continue
        col_idx = headers.index(header) + 1
        values: list[str | None] = []
        for row in range(2, n_rows + 2):
            fill = ws.cell(row=row, column=col_idx).fill
            has_rgb = fill.fill_type and fill.fgColor.type == "rgb"
            values.append(fill.fgColor.rgb if has_rgb else None)
        colors[col] = values
    return pd.DataFrame(colors, index=range(n_rows))


def load_upa_excel(
    path: Path | str = DEFAULT_EXCEL_PATH,
) -> pd.DataFrame:
    """Lê todas as abas do Excel e concatena pacientes.

    Args:
        path: Caminho ao arquivo Excel.

    Returns:
        DataFrame com colunas renomeadas para snake_case.
    """
    path = Path(path)
    xl = pd.ExcelFile(path)
    wb = openpyxl.load_workbook(path)
    frames: list[pd.DataFrame] = []

    for sheet in xl.sheet_names:
        df = pd.read_excel(path, sheet_name=sheet)
        colors = _fill_colors(wb[sheet], len(df))
        df = _harmonize_columns(df)
        df = _drop_junk_columns(df)
        df = pd.concat([df, colors], axis=1)
        df[COL_MES_COLETA] = sheet
        df[COL_LINHA_EXCEL] = df.index + 2
        frames.append(df)

    combined = pd.concat(frames, ignore_index=True)
    rename_map = {
        k: v for k, v in EXCEL_COL_MAP.items() if k in combined.columns
    }
    return combined.rename(columns=rename_map)
